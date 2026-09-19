"""Write a localized list snapshot without formulas or runtime secrets."""

from __future__ import annotations

import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_list_xlsx(
    path: Path,
    headers: list[str],
    rows: list[list[str]],
    sheet_title: str,
    *,
    overwrite: bool = False,
) -> None:
    """Stage a workbook beside its destination; preserve it if saving fails."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    try:
        for row_number, values in enumerate([headers, *rows], start=1):
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, value)
                # Keep leading zeroes and formula-like user descriptions as text.
                cell.data_type = "s"
                cell.number_format = "@"
                if row_number == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="305A80")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(headers) + 1):
            width = max(len(str(cell.value or "")) for cell in sheet[get_column_letter(column)])
            sheet.column_dimensions[get_column_letter(column)].width = min(60, max(12, width + 2))
        temporary_path = None
        try:
            with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            if overwrite:
                os.replace(temporary_path, path)
            else:
                # Exclusive publication also protects files created after the save dialog.
                os.link(temporary_path, path)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
    finally:
        workbook.close()
