from copy import deepcopy

import pytest

from config.settings import AppSettings, load_settings, save_settings
from config.terminology import normalize_terminology, validate_term
from gui.i18n import TRANSLATIONS, Translator
from gui.plain_text import message_text
from openvpn.addressing import ValidationError, calculate_addresses, validate_branch_number


def custom_terms(language="en", singular="Gateway", plural="Gateways"):
    return {language: {"router": {"singular": singular, "plural": plural}}}


@pytest.mark.parametrize("language,expected", [("en", ("Location", "Routers", "Device")), ("de", ("Standort", "Router", "Gerät"))])
def test_general_defaults(language, expected):
    t = Translator(language)
    assert (t.term("location"), t.term("router", "plural"), t.term("device")) == expected
    assert t("oven_ip") == f"IP ({expected[2]})"
    assert calculate_addresses("0004", 4, 1).name == "0004_Router1"
    assert set(TRANSLATIONS["en"]) == set(TRANSLATIONS["de"])


def test_terms_round_trip_and_language_fallback(tmp_path):
    path = tmp_path / "settings.json"
    terms = custom_terms(singular="  Gateway  ")
    terms["de"] = {"device": {"singular": "Maschine", "plural": ""}}
    save_settings(AppSettings(terminology=terms, vpn_server_host="vpn.example.com"), path)
    loaded = load_settings(path)
    t = Translator("en", loaded.terminology)
    assert t.term("router") == "Gateway"
    t.set_language("de")
    assert t.term("router") == "Router"
    assert t.term("device") == "Maschine"
    assert t.term("device", "plural") == "Geräte"
    assert loaded.vpn_server_host == "vpn.example.com"


@pytest.mark.parametrize("raw", [None, [], "bad", {"en": []}, {"en": {"router": 4}}, {"de": {"device": {"singular": 3, "plural": "x" * 41}}}])
def test_malformed_terms_fall_back(raw):
    assert normalize_terminology(raw) == {}


@pytest.mark.parametrize("value", ["x" * 41, "a\nb", "a\tb", "a\x00b", "a\u202eb"])
def test_invalid_terms_are_rejected(value):
    with pytest.raises(ValueError):
        validate_term(value)
    t = Translator("en", custom_terms(singular=value))
    assert t.term("router") == "Router"
    assert t.term("router", "plural") == "Gateways"


def test_custom_text_is_not_reinterpreted():
    term = "<b>{0} & @device_singular@</b>"
    t = Translator("en", custom_terms(singular=term, plural=term))
    assert t("delete_mango_question", "0004_Router1") == f"Delete {term} “0004_Router1”?"
    assert t("branch_device_count", 2, 3) == f"Locations: 2 / {term}: 3"
    assert t("xlsx_saved", "{0}@router_singular@") == "List saved to:\n{0}@router_singular@"
    assert message_text(term) == "<qt>&lt;b&gt;{0} &amp; @device_singular@&lt;/b&gt;</qt>"


def test_validation_uses_current_terms():
    t = Translator("en", {"en": {"location": {"singular": "Customer"}}})
    with pytest.raises(ValidationError) as result:
        validate_branch_number("../bad")
    assert "Customer" in t.error(result.value)
    assert "Branch" not in t.error(result.value)


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    application = QApplication.instance() or QApplication([])
    from pathlib import Path
    from PySide6.QtGui import QFontDatabase
    if "Segoe UI" not in QFontDatabase.families():
        for name in ("segoeui.ttf", "segoeuib.ttf"):
            font = Path("C:/Windows/Fonts") / name
            if font.exists():
                QFontDatabase.addApplicationFont(str(font))
    yield application


def test_dialog_drafts_cancel_apply_reset_and_languages(app):
    from PySide6.QtWidgets import QDialog
    from gui.terminology_dialog import TerminologyDialog
    original = custom_terms()
    before = deepcopy(original)
    dialog = TerminologyDialog(Translator("en", original), original)
    dialog.fields["router", "singular"].setText("Edge")
    dialog.language_combo.setCurrentIndex(dialog.language_combo.findData("de"))
    dialog.fields["device", "singular"].setText("Maschine")
    dialog.reset_language()
    assert dialog.fields["device", "singular"].text() == ""
    dialog.language_combo.setCurrentIndex(dialog.language_combo.findData("en"))
    assert dialog.fields["router", "singular"].text() == "Edge"
    dialog.reject()
    assert original == before
    assert dialog.result() == QDialog.Rejected
    dialog = TerminologyDialog(Translator("en"), original)
    dialog.fields["router", "singular"].setText("  Edge  ")
    dialog.accept()
    assert dialog.result() == QDialog.Accepted
    assert dialog.value()["en"]["router"]["singular"] == "Edge"
    assert original == before


def test_dialog_rejects_invalid_input_across_languages(app):
    from gui.terminology_dialog import TerminologyDialog
    dialog = TerminologyDialog(Translator(), {})
    dialog.fields["router", "singular"].setText("x" * 41)
    assert not dialog.apply_button.isEnabled()
    dialog.language_combo.setCurrentIndex(1)
    assert not dialog.apply_button.isEnabled()
    dialog.language_combo.setCurrentIndex(0)
    dialog.reset_language()
    assert dialog.apply_button.isEnabled()
    dialog.fields["device", "singular"].setText("<b>{0}</b>")
    assert "IP (<b>{0}</b>)" in dialog.preview_label.text()
    from PySide6.QtCore import Qt
    assert dialog.preview_label.textFormat() == Qt.PlainText
    dialog.close()


@pytest.fixture
def window(app, tmp_path, monkeypatch):
    from database.database import Database
    from gui import main_window, main_window_view
    from openvpn.runtime import OpenVPNRuntimePaths, StatusSnapshot
    monkeypatch.setattr(main_window, "discover_openvpn_runtime", lambda _: OpenVPNRuntimePaths(tmp_path / "config", tmp_path / "log"))
    monkeypatch.setattr(main_window, "openvpn_service_running", lambda _: False)
    monkeypatch.setattr(main_window, "read_status_snapshot", lambda _: StatusSnapshot(False, False, None, {}))
    monkeypatch.setattr(main_window_view, "save_settings", lambda settings: save_settings(settings, tmp_path / "settings.json"))
    db = Database(tmp_path / "test.db")
    db.add_branch_with_mangos("0004", 2)
    db.add_branch_with_mangos("0005", 1)
    settings = AppSettings(pki_path=tmp_path / "pki", openvpn_root=tmp_path / "OpenVPN", easyrsa_root=tmp_path / "easy-rsa")
    widget = main_window.MainWindow(db, settings)
    widget.status_timer.stop()
    yield widget
    widget.close()
    db.close()


def test_apply_preserves_data_selection_expansion_and_export(window, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from openpyxl import load_workbook
    branch = window.tree.topLevelItem(0)
    branch.setExpanded(True)
    window.tree.setCurrentItem(branch.child(0))
    selected = window._selected_mango_id()
    window.database.set_certificate_created(selected, True)
    window.database.mark_configs_created([selected])
    before = list(window.database.connection.iterdump())
    changes = window.database.connection.total_changes
    terms = custom_terms(singular="<Edge>{0}", plural="Gateways")
    terms["en"]["location"] = {"plural": "L" * 40 + ""}
    terms["en"]["device"] = {"singular": "Machine"}
    assert window.apply_terminology(terms)
    assert list(window.database.connection.iterdump()) == before
    assert window.database.connection.total_changes == changes
    assert window._selected_mango_id() == selected
    assert branch.isExpanded()
    assert not window.tree.topLevelItem(1).isExpanded()
    assert window.add_mango_button.text() == "Add: <Edge>{0}"
    assert window.device_table.headerItem().text(6) == "IP (<Edge>{0})"
    assert window.device_table.headerItem().text(7) == "IP (Machine)"
    assert "Gateways" in window.workflow_summary_label.text()
    assert "Gateways: 3" in window.count_status_label.text()
    target = tmp_path / "list.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    monkeypatch.setattr(QMessageBox, "information", lambda *a: QMessageBox.Ok)
    window.export_list()
    workbook = load_workbook(target)
    assert workbook.active.cell(1, 5).value == "IP (<Edge>{0})"
    assert workbook.active.cell(1, 6).value == "IP (Machine)"
    assert workbook.active.max_row == 4
    assert len(workbook.active.title) <= 31
    workbook.close()
    assert list(window.database.connection.iterdump()) == before
    assert load_settings(tmp_path / "settings.json").terminology == terms


def test_failed_save_does_not_apply(window, monkeypatch):
    from gui import main_window_view
    from PySide6.QtWidgets import QMessageBox
    def fail(_):
        raise OSError("Read-only settings")
    monkeypatch.setattr(main_window_view, "save_settings", fail)
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: QMessageBox.Ok)
    assert not window.apply_terminology(custom_terms())
    assert window.settings.terminology == {}
    assert window.t.term("router") == "Router"


@pytest.mark.parametrize("title", ["'[]:*?/\\'", "a" * 90, "History", "", "A" * 30 + "'B"])
def test_excel_titles_are_valid_and_headers_remain_literal(tmp_path, title):
    from openpyxl import load_workbook
    from gui.spreadsheet import write_list_xlsx
    path = tmp_path / "terms.xlsx"
    write_list_xlsx(path, ["=1+1", "<b>{0}</b>"], [["0004", "data"]], title)
    workbook = load_workbook(path)
    assert 1 <= len(workbook.active.title) <= 31
    assert not any(c in workbook.active.title for c in "[]:*?/\\")
    assert workbook.active.cell(1, 1).data_type == "s"
    assert workbook.active.cell(1, 1).value == "=1+1"
    assert workbook.active.cell(1, 2).value == "<b>{0}</b>"
    workbook.close()


@pytest.mark.parametrize("language", ["en", "de"])
@pytest.mark.parametrize("theme", ["light", "dark"])
def test_dialog_layout_and_plain_labels(app, window, language, theme):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QLabel
    from gui.terminology_dialog import TerminologyDialog
    from gui.branch_dialog import BranchDialog
    from gui.mango_dialog import MangoDialog
    from gui.certificate_dialog import CertificateDialog
    from gui.export_dialog import ExportDialog
    from gui.theme import apply_theme
    terms = {language: {role: {form: "<b>{0} Long term " + "X" * 20 for form in ("singular", "plural")} for role in ("location", "router", "device")}}
    t = Translator(language, terms)
    apply_theme(app, theme)
    dialogs = [
        TerminologyDialog(t, terms),
        BranchDialog(t),
        MangoDialog(t, window.database.list_branches()),
        CertificateDialog(t, window.database, window.settings),
        ExportDialog(t, window.database, window.settings, None, None),
    ]
    try:
        for dialog in dialogs:
            dialog.show()
            app.processEvents()
            assert dialog.width() <= dialog.screen().availableGeometry().width()
            for label in dialog.findChildren(QLabel):
                if "<b>" in label.text():
                    assert label.textFormat() == Qt.PlainText
        certificate_dialog = dialogs[3]
        assert "<b>{0}" in certificate_dialog.create_all_button.text()
        button = certificate_dialog.create_all_button
        assert button.height() >= button.heightForWidth(button.width())
    finally:
        for dialog in dialogs:
            dialog.hide()
            dialog.deleteLater()
        app.processEvents()


@pytest.mark.parametrize("theme", ["light", "dark"])
def test_terminology_dialog_content_uses_theme_surface(app, theme):
    from PySide6.QtGui import QColor
    from PySide6.QtWidgets import QWidget

    from gui.terminology_dialog import TerminologyDialog
    from gui.theme import apply_theme
    from gui.theme_palette import theme_color

    apply_theme(app, theme)
    dialog = TerminologyDialog(Translator("en"), {})
    dialog.show()
    app.processEvents()
    content = dialog.findChild(QWidget, "terminologyContent")
    expected = QColor(theme_color(theme, "surface"))
    assert content.palette().color(content.backgroundRole()) == expected
    dialog.close()
