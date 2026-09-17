from pathlib import Path

import pytest
from openpyxl import load_workbook

from gui.spreadsheet import write_list_xlsx


def test_xlsx_keeps_leading_zeroes_and_formula_like_text(tmp_path):
    path = tmp_path / "list.xlsx"
    write_list_xlsx(path, ["Branch", "Description"], [["0004", '=HYPERLINK("bad")']], "List")
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet["A2"].value == "0004"
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("bad")'
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B2"
    workbook.close()


def test_xlsx_refuses_unconfirmed_overwrite_and_cleans_staging(tmp_path):
    path = tmp_path / "list.xlsx"
    path.write_bytes(b"existing file")
    with pytest.raises(FileExistsError):
        write_list_xlsx(path, ["Name"], [], "List")
    assert path.read_bytes() == b"existing file"
    assert list(tmp_path.iterdir()) == [path]
    write_list_xlsx(path, ["Name"], [], "List", overwrite=True)
    workbook = load_workbook(path)
    assert workbook.active["A1"].value == "Name"
    workbook.close()


def test_xlsx_failed_save_preserves_existing_file(tmp_path, monkeypatch):
    path = tmp_path / "list.xlsx"
    path.write_bytes(b"existing file")

    def fail_save(self, target):
        Path(target).write_bytes(b"partial")
        raise OSError("Disk full")

    monkeypatch.setattr("gui.spreadsheet.Workbook.save", fail_save)
    with pytest.raises(OSError, match="Disk full"):
        write_list_xlsx(path, ["Name"], [], "List", overwrite=True)
    assert path.read_bytes() == b"existing file"
    assert list(tmp_path.iterdir()) == [path]
