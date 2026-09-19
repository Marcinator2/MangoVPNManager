from types import SimpleNamespace

import pytest

from gui.main_window import (
    MainWindow,
    NORMAL_STATUS_INTERVAL_MS,
    STARTUP_STATUS_POLL_COUNT,
)


class FakeTimer:
    def __init__(self) -> None:
        self.intervals: list[int] = []

    def setInterval(self, interval: int) -> None:
        self.intervals.append(interval)


def test_runtime_poll_switches_to_normal_interval_after_ten_startup_checks() -> None:
    refreshes: list[bool] = []
    window = SimpleNamespace(
        _startup_status_checks_remaining=STARTUP_STATUS_POLL_COUNT,
        status_timer=FakeTimer(),
        _refresh_runtime_status=lambda: refreshes.append(True),
    )

    for _ in range(STARTUP_STATUS_POLL_COUNT - 1):
        MainWindow._poll_runtime_status(window)

    assert len(refreshes) == 9
    assert window.status_timer.intervals == []

    MainWindow._poll_runtime_status(window)

    assert len(refreshes) == 10
    assert window._startup_status_checks_remaining == 0
    assert window.status_timer.intervals == [NORMAL_STATUS_INTERVAL_MS]


@pytest.mark.parametrize("language", ["en", "de"])
def test_oven_settings_and_complete_list_export(tmp_path, monkeypatch, language):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
    from openpyxl import load_workbook
    from config.settings import AppSettings
    from database.database import Database
    from gui import main_window
    from openvpn.runtime import OpenVPNRuntimePaths, StatusSnapshot

    application = QApplication.instance() or QApplication([])
    monkeypatch.setattr(main_window, "discover_openvpn_runtime", lambda _: OpenVPNRuntimePaths(tmp_path / "config", tmp_path / "log"))
    monkeypatch.setattr(main_window, "openvpn_service_running", lambda _: False)
    monkeypatch.setattr(main_window, "read_status_snapshot", lambda _: StatusSnapshot(False, False, None, {}))
    database = Database(tmp_path / "test.db")
    database.add_branch_with_mangos("0004", 2, "Test branch")
    database.add_branch_with_mangos("0005", 0, "Empty branch")
    database.add_branch_with_mangos("0006", 1, "Another branch")
    window = MainWindow(database, AppSettings(language=language, pki_path=tmp_path / "pki"))
    try:
        window.show()
        application.processEvents()
        branch = window.tree.topLevelItem(0)
        assert not branch.isExpanded()
        assert branch.text(0) == "(2) 0004"
        assert window.tree.topLevelItem(1).text(0) == "(0) 0005"
        assert not window.tree.topLevelItem(1).isExpanded()
        branch.setExpanded(True)
        window.reload_tree()
        branch = window.tree.topLevelItem(0)
        assert branch.isExpanded()
        branch.setExpanded(False)
        window.reload_tree()
        branch = window.tree.topLevelItem(0)
        assert not branch.isExpanded()
        child = branch.child(1)
        assert [child.text(c) for c in (5, 6, 7)] == ["10.1.2.102", "255.255.255.0", "10.1.2.1"]
        assert child.text(8) == window.t("no")
        assert child.text(9) == window.t("no")
        assert window.t("unknown") in child.text(10)
        assert branch.text(1) == "Test branch"
        assert window.tree.headerItem().text(1) == window.t("description")
        assert window.device_table.topLevelItemCount() == 3
        assert window.device_table.isColumnHidden(3)
        window.tree.setCurrentItem(branch)
        application.processEvents()
        assert window.device_table.topLevelItemCount() == 2
        device = window.device_table.topLevelItem(0)
        window.device_table.setCurrentItem(device)
        application.processEvents()
        selected_id = device.data(0, main_window.ROLE_ID)
        assert window._selected_mango_id() == selected_id
        selected_mango = window._mango(selected_id)
        assert selected_mango is not None
        assert device.text(1) == selected_mango.branch_number
        assert device.text(3) == "1"
        assert device.text(6) == selected_mango.mango_ip
        assert device.text(7) == selected_mango.oven_ip
        assert device.text(8) == selected_mango.oven_subnet_mask
        assert device.text(9) == selected_mango.oven_gateway
        assert device.text(10) == window.t("no")
        assert device.text(11) == window.t("no")
        assert window.t("unknown") in device.text(12)
        window._select_all_mangos()
        assert window.device_table.topLevelItemCount() == 3
        path = tmp_path / "list.xlsx"
        monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(path.with_suffix("")), ""))
        monkeypatch.setattr(QMessageBox, "information", lambda *args: QMessageBox.Ok)
        window.export_list()
        workbook = load_workbook(path)
        sheet = workbook.active
        assert sheet.max_row == 4
        assert sheet.max_column == 14
        assert sheet["A2"].value == "0004_Mango1"
        assert sheet["A3"].value == "0004_Mango2"
        assert sheet["B1"].value == window.t("description")
        assert sheet["B2"].value == "Test branch"
        assert sheet["B3"].value == "Test branch"
        assert sheet["A4"].value == "0006_Mango1"
        assert sheet["B4"].value == "Another branch"
        assert sheet["G1"].value == window.t("oven_subnet_mask")
        assert sheet["H3"].value == "10.1.2.1"
        workbook.close()
        original = path.read_bytes()
        monkeypatch.setattr(QMessageBox, "question", lambda *args: QMessageBox.No)
        window.export_list()
        assert path.read_bytes() == original
    finally:
        window.close()
